"""
Build Millionaires Tax legislative-district geographic weights v0.1.

Purpose
-------
Extract the existing Model 2 estimated tax-paid distribution from the
project's source workbook and convert it into geographic incidence shares.

The original Model 2 dollar estimates are NOT treated as the statewide
FY2029 or FY2030 revenue forecast.

Instead:

    Model 2 -> geographic distribution
    DOR      -> statewide scenario-year revenue control

For each legislative district:

    tax_share_d =
        original_model2_tax_paid_d / statewide_original_model2_tax_paid

    FY2029 tax incidence =
        tax_share_d * $2.698B

    FY2030 tax incidence =
        tax_share_d * $3.732B

This preserves the original geographic model while normalizing its
statewide total to the official DOR scenario-year controls.
"""

from pathlib import Path
import csv

from openpyxl import load_workbook


# ---------------------------------------------------------------------------
# Paths
# ---------------------------------------------------------------------------

INPUT_FILE = Path(
    "Washington State Millionaires Tax by Legislative "
    "District_currency_cleaned NEWEST.xlsx"
)

OUTPUT_DIR = Path("data/processed/model")
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

OUTPUT_FILE = (
    OUTPUT_DIR / "millionaires_tax_ld_weights_v0_1.csv"
)


# ---------------------------------------------------------------------------
# Controls
# ---------------------------------------------------------------------------

EXPECTED_LD_COUNT = 49

FY2029_DOR_CONTROL = 2_698_000_000
FY2030_DOR_CONTROL = 3_732_000_000


def clean_header(value):
    if value is None:
        return ""
    return str(value).strip()


def find_header_row_and_columns(ws):
    """
    Locate the row containing the required source headers.

    Required logical fields:
        Legislative District
        Tax Paid

    The workbook previously used Column P for Model 2 Tax Paid, but this
    function searches headers rather than relying only on a column letter.
    """

    for row_number in range(1, min(ws.max_row, 25) + 1):

        headers = {
            clean_header(ws.cell(row_number, col).value): col
            for col in range(1, ws.max_column + 1)
            if ws.cell(row_number, col).value is not None
        }

        if (
            "Legislative District" in headers
            and "Tax Paid" in headers
        ):
            return (
                row_number,
                headers["Legislative District"],
                headers["Tax Paid"],
            )

    raise ValueError(
        "Could not locate a header row containing both "
        "'Legislative District' and 'Tax Paid'."
    )


def normalize_district(value):
    """
    Convert a legislative-district cell to integer 1-49.

    Accepts numeric cells and simple strings such as:
        1
        1.0
        LD 1
        District 1
    """

    if value is None:
        return None

    if isinstance(value, (int, float)):
        district = int(value)
        if 1 <= district <= 49:
            return district
        return None

    text = str(value).strip().lower()

    text = text.replace("legislative district", "")
    text = text.replace("district", "")
    text = text.replace("ld", "")
    text = text.strip()

    try:
        district = int(float(text))
    except ValueError:
        return None

    if 1 <= district <= 49:
        return district

    return None


def normalize_money(value):
    """
    Convert Excel numeric/currency values to float.
    """

    if value is None:
        return None

    if isinstance(value, (int, float)):
        return float(value)

    text = str(value).strip()

    if not text:
        return None

    text = (
        text.replace("$", "")
        .replace(",", "")
        .replace("(", "-")
        .replace(")", "")
    )

    return float(text)


def inspect_workbook():
    """
    Find the worksheet containing the required Model 2 fields.
    """

    if not INPUT_FILE.exists():
        raise FileNotFoundError(
            f"Workbook not found: {INPUT_FILE}"
        )

    wb = load_workbook(
        INPUT_FILE,
        data_only=True,
        read_only=True,
    )

    candidates = []

    for ws in wb.worksheets:
        try:
            header_row, ld_col, tax_col = (
                find_header_row_and_columns(ws)
            )

            candidates.append(
                {
                    "worksheet": ws.title,
                    "header_row": header_row,
                    "ld_col": ld_col,
                    "tax_col": tax_col,
                }
            )

        except ValueError:
            continue

    if not candidates:
        raise ValueError(
            "No worksheet contains both required headers."
        )

    if len(candidates) > 1:
        print()
        print("WARNING")
        print("-" * 80)
        print(
            "More than one worksheet contains the required headers:"
        )
        for candidate in candidates:
            print(
                f"  {candidate['worksheet']} "
                f"(header row {candidate['header_row']})"
            )
        print(
            "Using the first matching worksheet."
        )

    return wb, candidates[0]


def extract_rows(wb, source):
    ws = wb[source["worksheet"]]

    rows_by_district = {}

    for row_number in range(
        source["header_row"] + 1,
        ws.max_row + 1,
    ):
        district = normalize_district(
            ws.cell(
                row_number,
                source["ld_col"],
            ).value
        )

        if district is None:
            continue

        tax_paid = normalize_money(
            ws.cell(
                row_number,
                source["tax_col"],
            ).value
        )

        if tax_paid is None:
            raise ValueError(
                f"LD {district} has no Model 2 Tax Paid value "
                f"at worksheet row {row_number}."
            )

        if district in rows_by_district:
            raise ValueError(
                f"Duplicate Legislative District {district}."
            )

        rows_by_district[district] = {
            "legislative_district": district,
            "original_model2_tax_paid": tax_paid,
            "source_worksheet": source["worksheet"],
            "source_row": row_number,
        }

    rows = [
        rows_by_district[d]
        for d in sorted(rows_by_district)
    ]

    return rows


def build_weights(rows):
    statewide_original = sum(
        row["original_model2_tax_paid"]
        for row in rows
    )

    if statewide_original <= 0:
        raise ValueError(
            "Original Model 2 statewide tax total must be positive."
        )

    for row in rows:
        share = (
            row["original_model2_tax_paid"]
            / statewide_original
        )

        row["tax_incidence_share"] = share

        row["tax_paid_fy2029_2_698b"] = (
            share * FY2029_DOR_CONTROL
        )

        row["tax_paid_fy2030_3_732b"] = (
            share * FY2030_DOR_CONTROL
        )

    return statewide_original


def validate(rows, statewide_original):
    assert len(rows) == EXPECTED_LD_COUNT, (
        f"Expected {EXPECTED_LD_COUNT} LDs, "
        f"found {len(rows)}."
    )

    districts = [
        row["legislative_district"]
        for row in rows
    ]

    assert districts == list(range(1, 50)), (
        "Legislative districts are not exactly 1-49."
    )

    assert statewide_original > 0

    share_sum = sum(
        row["tax_incidence_share"]
        for row in rows
    )

    assert abs(share_sum - 1.0) < 1e-12, (
        f"Tax-incidence shares sum to {share_sum}."
    )

    fy2029_sum = sum(
        row["tax_paid_fy2029_2_698b"]
        for row in rows
    )

    fy2030_sum = sum(
        row["tax_paid_fy2030_3_732b"]
        for row in rows
    )

    assert abs(
        fy2029_sum - FY2029_DOR_CONTROL
    ) < 0.01

    assert abs(
        fy2030_sum - FY2030_DOR_CONTROL
    ) < 0.01


def write_csv(rows):
    fieldnames = [
        "legislative_district",
        "original_model2_tax_paid",
        "tax_incidence_share",
        "tax_paid_fy2029_2_698b",
        "tax_paid_fy2030_3_732b",
        "source_worksheet",
        "source_row",
    ]

    with OUTPUT_FILE.open(
        "w",
        newline="",
        encoding="utf-8",
    ) as handle:

        writer = csv.DictWriter(
            handle,
            fieldnames=fieldnames,
        )

        writer.writeheader()
        writer.writerows(rows)


def money(value):
    return f"${value:,.2f}"


def print_report(
    rows,
    statewide_original,
    source,
):
    print()
    print(
        "MILLIONAIRES TAX LD GEOGRAPHIC WEIGHTS v0.1"
    )
    print("=" * 80)

    print()
    print("SOURCE")
    print("-" * 80)
    print(f"Workbook            : {INPUT_FILE}")
    print(
        f"Worksheet           : "
        f"{source['worksheet']}"
    )
    print(
        f"Header row          : "
        f"{source['header_row']}"
    )
    print(
        f"LD column           : "
        f"{source['ld_col']}"
    )
    print(
        f"Tax Paid column     : "
        f"{source['tax_col']}"
    )

    print()
    print("ORIGINAL MODEL 2")
    print("-" * 80)
    print(
        f"Legislative districts          : "
        f"{len(rows)}"
    )
    print(
        f"Original Model 2 statewide tax : "
        f"{money(statewide_original)}"
    )
    print(
        f"Tax-incidence shares sum       : "
        f"{sum(r['tax_incidence_share'] for r in rows):.12f}"
    )

    print()
    print("NORMALIZED DOR CONTROLS")
    print("-" * 80)

    fy2029_total = sum(
        r["tax_paid_fy2029_2_698b"]
        for r in rows
    )

    fy2030_total = sum(
        r["tax_paid_fy2030_3_732b"]
        for r in rows
    )

    print(
        f"FY2029 normalized tax total    : "
        f"{money(fy2029_total)}"
    )
    print(
        f"FY2030 normalized tax total    : "
        f"{money(fy2030_total)}"
    )

    print()
    print("LARGEST MODELED TAX-INCIDENCE SHARES")
    print("-" * 80)

    ranked = sorted(
        rows,
        key=lambda row: row["tax_incidence_share"],
        reverse=True,
    )

    for row in ranked[:10]:
        print(
            f"LD {row['legislative_district']:>2}: "
            f"original={money(row['original_model2_tax_paid']):>18}  "
            f"share={row['tax_incidence_share']:>7.3%}  "
            f"FY2029={money(row['tax_paid_fy2029_2_698b'])}"
        )

    print()
    print("IMPORTANT")
    print("-" * 80)
    print(
        "The original Model 2 dollar values are used only to derive "
        "the geographic distribution of estimated tax incidence."
    )
    print(
        "They are not treated as the FY2029 or FY2030 statewide "
        "Millionaires Tax revenue forecast."
    )
    print(
        "The geographic shares are normalized to the official DOR "
        "scenario-year statewide controls."
    )

    print()
    print("OUTPUT")
    print("-" * 80)
    print(OUTPUT_FILE)


def main():
    wb, source = inspect_workbook()

    rows = extract_rows(
        wb,
        source,
    )

    statewide_original = build_weights(rows)

    validate(
        rows,
        statewide_original,
    )

    write_csv(rows)

    print_report(
        rows,
        statewide_original,
        source,
    )


if __name__ == "__main__":
    main()