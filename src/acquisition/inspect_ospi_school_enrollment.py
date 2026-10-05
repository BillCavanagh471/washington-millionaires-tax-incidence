from pathlib import Path

import openpyxl


SOURCE = Path(
    "data/raw/ospi/enrollment/"
    "schoolenrollment-17-18-current.xlsx"
)


def main():

    if not SOURCE.exists():
        raise FileNotFoundError(
            f"Source file not found: {SOURCE}"
        )

    print("OSPI SCHOOL ENROLLMENT WORKBOOK PROFILE")
    print("=" * 80)
    print(f"Source: {SOURCE}")
    print(
        f"Size  : {SOURCE.stat().st_size:,} bytes"
    )
    print()

    workbook = openpyxl.load_workbook(
        SOURCE,
        read_only=True,
        data_only=True,
    )

    print("WORKSHEETS")
    print("-" * 80)

    for i, name in enumerate(
        workbook.sheetnames,
        start=1,
    ):
        ws = workbook[name]

        print(
            f"{i:>2}. {name!r}"
            f"  rows={ws.max_row:,}"
            f"  columns={ws.max_column:,}"
        )

    print()
    print("FIRST 12 ROWS OF EACH WORKSHEET")
    print("=" * 80)

    for name in workbook.sheetnames:

        ws = workbook[name]

        print()
        print(f"SHEET: {name}")
        print("-" * 80)

        for row_number, row in enumerate(
            ws.iter_rows(
                min_row=1,
                max_row=min(ws.max_row, 12),
                values_only=True,
            ),
            start=1,
        ):

            values = [
                "" if value is None else str(value)
                for value in row
            ]

            print(
                f"{row_number:>3}: "
                + " | ".join(values)
            )

    workbook.close()


if __name__ == "__main__":
    main()